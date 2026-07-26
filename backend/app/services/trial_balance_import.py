"""Parsing et validation des imports de balance de vérification (1 ou N mois)."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO, StringIO
import csv
from typing import Iterable

from openpyxl import load_workbook

BALANCE_SUM_TOLERANCE = 0.01
MAX_MULTI_MONTH_COUNT = 120

MONTHS_FR = [
    "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
    "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre",
]


@dataclass(frozen=True)
class MonthPeriod:
    year: int
    month: int  # 1–12


@dataclass
class SingleMonthTrialBalanceRow:
    account_number: str
    description: str
    amount: float


@dataclass
class MultiMonthTrialBalanceRow:
    account_number: str
    description: str
    amounts: list[float]


@dataclass
class MultiMonthParseResult:
    rows: list[MultiMonthTrialBalanceRow]
    periods: list[MonthPeriod]
    month_sums: list[float]


def excel_account_number_str(raw) -> str:
    if raw is None:
        return ""
    if isinstance(raw, bool):
        return str(raw)
    if isinstance(raw, int):
        return str(raw)
    if isinstance(raw, float):
        if raw.is_integer():
            return str(int(raw))
        return str(raw).strip()
    s = str(raw).strip()
    if len(s) > 2 and s.endswith(".0") and s[:-2].replace("-", "").isdigit():
        return s[:-2]
    return s


def parse_amount_cell(raw) -> float:
    if raw is None or raw == "":
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).replace(",", ".").replace(" ", "").replace("\u00a0", "").strip()
    if not s:
        return 0.0
    return float(s)


def calendar_months_from_start(start_year: int, start_month: int, count: int) -> list[MonthPeriod]:
    if count < 1 or count > MAX_MULTI_MONTH_COUNT:
        raise ValueError(f"month_count must be between 1 and {MAX_MULTI_MONTH_COUNT}")
    if start_month < 1 or start_month > 12:
        raise ValueError("start_month must be between 1 and 12")

    periods: list[MonthPeriod] = []
    year, month = start_year, start_month
    for _ in range(count):
        periods.append(MonthPeriod(year, month))
        month += 1
        if month > 12:
            month = 1
            year += 1
    return periods


def period_label(period: MonthPeriod) -> str:
    return f"{MONTHS_FR[period.month - 1]} {period.year}"


def flatten_multi_month_for_unknown_preview(
    rows: Iterable[MultiMonthTrialBalanceRow],
) -> list[tuple[str, str, float]]:
    out: list[tuple[str, str, float]] = []
    for row in rows:
        for amount in row.amounts:
            out.append((row.account_number, row.description, amount))
    return out


def _parse_single_month_matrix(matrix: list[list]) -> list[SingleMonthTrialBalanceRow]:
    rows_out: list[SingleMonthTrialBalanceRow] = []
    for row in matrix[1:]:
        if not row or len(row) < 3 or row[0] is None:
            continue
        account_number = excel_account_number_str(row[0])
        if not account_number:
            continue
        description = str(row[1] or "").strip() if len(row) > 1 else ""
        try:
            amount = parse_amount_cell(row[2] if len(row) > 2 else 0)
        except (ValueError, TypeError):
            continue
        rows_out.append(SingleMonthTrialBalanceRow(account_number, description, amount))
    return rows_out


def parse_single_month_file(content: bytes, filename: str | None) -> tuple[list[SingleMonthTrialBalanceRow], float]:
    matrix = _load_file_matrix(content, filename)
    rows = _parse_single_month_matrix(matrix)
    total = sum(r.amount for r in rows)
    return rows, total


def _load_file_matrix(content: bytes, filename: str | None) -> list[list]:
    fn = (filename or "").lower()
    if fn.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return list(csv.reader(StringIO(text)))
    wb = load_workbook(BytesIO(content), read_only=True)
    ws = wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _parse_multi_month_matrix(
    matrix: list[list],
    month_count: int,
) -> list[MultiMonthTrialBalanceRow]:
    if not matrix:
        return []

    required_cols = 2 + month_count
    max_data_cols = 0
    for row in matrix[1:]:
        if row and row[0] is not None and excel_account_number_str(row[0]):
            max_data_cols = max(max_data_cols, len(row))

    if max_data_cols < required_cols:
        raise ValueError(
            f"Le fichier doit contenir au moins {month_count} colonne(s) de montants "
            f"(colonnes C à {chr(ord('C') + month_count - 1)}), "
            f"mais seulement {max(0, max_data_cols - 2)} trouvée(s)."
        )

    rows_out: list[MultiMonthTrialBalanceRow] = []
    for row in matrix[1:]:
        if not row or len(row) < 2 or row[0] is None:
            continue
        account_number = excel_account_number_str(row[0])
        if not account_number:
            continue
        description = str(row[1] or "").strip()
        amounts: list[float] = []
        for col_idx in range(month_count):
            cell = row[2 + col_idx] if len(row) > 2 + col_idx else None
            try:
                amounts.append(parse_amount_cell(cell))
            except (ValueError, TypeError):
                amounts.append(0.0)
        rows_out.append(MultiMonthTrialBalanceRow(account_number, description, amounts))
    return rows_out


def parse_multi_month_file(
    content: bytes,
    filename: str | None,
    start_year: int,
    start_month: int,
    month_count: int,
) -> MultiMonthParseResult:
    periods = calendar_months_from_start(start_year, start_month, month_count)
    matrix = _load_file_matrix(content, filename)
    rows = _parse_multi_month_matrix(matrix, month_count)
    month_sums = [0.0] * month_count
    for row in rows:
        for idx, amount in enumerate(row.amounts):
            month_sums[idx] += amount
    return MultiMonthParseResult(rows=rows, periods=periods, month_sums=month_sums)


def month_sums_are_balanced(month_sums: Iterable[float]) -> bool:
    return all(abs(s) <= BALANCE_SUM_TOLERANCE for s in month_sums)
